from __future__ import annotations

import csv
import datetime as dt
import math
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

try:
    from openpyxl import load_workbook
except ImportError:  # CSV support remains available without the optional dependency.
    load_workbook = None


OUTPUT_COLUMNS = [
    "Financial Year",
    "year",
    "month",
    "day",
    "period",
    "projected demand",
]
ROOFTOP_OUTPUT_COLUMNS = [
    "Financial Year",
    "year",
    "month",
    "day",
    "period",
    "demand before rooftop",
    "incremental rooftop capacity",
    "incremental rooftop generation",
    "projected demand",
]
DATE_FORMATS = ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%y")


class ProfileGenerationError(ValueError):
    """Raised when inputs cannot produce a deterministic projection."""


@dataclass(frozen=True)
class BaseDay:
    date: dt.date
    values: tuple[float, ...]


@dataclass(frozen=True)
class YearSummary:
    projection_year: int
    row_count: int
    day_count: int
    periods_per_day: int
    base_profile_peak_mw: float
    base_profile_energy_gwh: float
    base_profile_is_normalized: bool
    target_peak_mw: float
    achieved_peak_mw: float
    target_energy_gwh: float
    achieved_energy_gwh: float
    mapped_base_peak_mw: float
    mapped_base_energy_gwh: float
    peak_growth_rate: float
    low_rank_growth_rate: float
    rooftop_enabled: bool = False
    before_rooftop_peak_mw: float | None = None
    before_rooftop_energy_gwh: float | None = None
    peak_reduction_mw: float | None = None
    peak_reduction_percent: float | None = None
    rooftop_generation_gwh: float | None = None
    rooftop_template_cuf_percent: float | None = None
    rooftop_effective_cuf_percent: float | None = None
    rooftop_year_end_capacity_mw: float | None = None
    adjusted_minimum_mw: float | None = None
    adjusted_peak_date: str | None = None
    adjusted_peak_period: int | None = None


@dataclass(frozen=True)
class RooftopProfile:
    mode: str
    periods_per_day: int
    values: dict[tuple[int, ...], float]
    template_cuf: float
    row_count: int


@dataclass(frozen=True)
class GenerationResult:
    output_path: Path
    rows_written: int
    summaries: tuple[YearSummary, ...]
    graph: dict


def parse_date(value: str) -> dt.date:
    value = str(value).strip()
    for fmt in DATE_FORMATS:
        try:
            return dt.datetime.strptime(value, fmt).date()
        except ValueError:
            pass
    raise ProfileGenerationError(
        f"Could not parse date '{value}'. Use YYYY-MM-DD or DD-MM-YYYY."
    )


def fiscal_year_label(start_year: int) -> str:
    return f"{start_year}-{str((start_year + 1) % 100).zfill(2)}"


def validate_base_financial_year_period(start: dt.date, end: dt.date) -> None:
    expected_end = dt.date(start.year + 1, 3, 31)
    if (start.month, start.day) != (4, 1) or end != expected_end:
        raise ProfileGenerationError(
            "Base period must run from 1 April through 31 March of the following year."
        )


def date_range(start: dt.date, end: dt.date) -> list[dt.date]:
    if end < start:
        raise ProfileGenerationError("End date must be on or after start date.")
    days = []
    day = start
    while day <= end:
        days.append(day)
        day += dt.timedelta(days=1)
    return days


def _clean_header(header: str | None) -> str:
    return "" if header is None else str(header).strip()


def _normal_header(header: str | None) -> str:
    return re.sub(r"[^a-z0-9]+", "", _clean_header(header).lower())


def _as_float(value: str, field_name: str) -> float:
    text = str(value).strip().replace(",", "")
    if not text:
        raise ProfileGenerationError(f"Missing numeric value in {field_name}.")
    try:
        return float(text)
    except ValueError as exc:
        raise ProfileGenerationError(
            f"Could not parse numeric value '{value}' in {field_name}."
        ) from exc


def _as_int(value: str, field_name: str) -> int:
    number = _as_float(value, field_name)
    if not number.is_integer():
        raise ProfileGenerationError(f"{field_name} must be an integer.")
    return int(number)


def _read_csv_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        if not reader.fieldnames:
            raise ProfileGenerationError(f"{path.name} has no header row.")
        fieldnames = [_clean_header(name) for name in reader.fieldnames]
        rows = []
        for row in reader:
            cleaned = {}
            for raw_key, value in row.items():
                key = _clean_header(raw_key)
                if key:
                    cleaned[key] = "" if value is None else str(value).strip()
            if any(value for value in cleaned.values()):
                rows.append(cleaned)
    return fieldnames, rows


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    return str(value).strip()


def _read_xlsx_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    if load_workbook is None:
        raise ProfileGenerationError(
            "XLSX support is unavailable. Install the dependencies from requirements.txt."
        )
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ProfileGenerationError(
            f"Could not read Excel workbook {path.name}: {exc}"
        ) from exc

    try:
        preferred = workbook.active
        worksheets = [preferred] + [
            sheet for sheet in workbook.worksheets if sheet is not preferred
        ]
        for worksheet in worksheets:
            values = worksheet.iter_rows(values_only=True)
            try:
                header_row = next(values)
            except StopIteration:
                continue
            fieldnames = [_clean_header(_cell_text(value)) for value in header_row]
            if not any(fieldnames):
                continue

            rows = []
            for row_values in values:
                cleaned = {}
                for index, name in enumerate(fieldnames):
                    if not name:
                        continue
                    value = row_values[index] if index < len(row_values) else None
                    cleaned[name] = _cell_text(value)
                if any(value for value in cleaned.values()):
                    rows.append(cleaned)
            if rows:
                return fieldnames, rows
    finally:
        workbook.close()

    raise ProfileGenerationError(f"{path.name} has no populated worksheet.")


def _read_tabular_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _read_csv_rows(path)
    if suffix == ".xlsx":
        return _read_xlsx_rows(path)
    raise ProfileGenerationError(f"{path.name} must be a CSV or XLSX file.")


def _find_required_column(fieldnames: Iterable[str], wanted: str, path: Path) -> str:
    for name in fieldnames:
        if _normal_header(name) == wanted.lower():
            return name
    raise ProfileGenerationError(f"{path.name} must contain a '{wanted}' column.")


def _find_value_column(
    fieldnames: Iterable[str],
    rows: list[dict[str, str]],
    excluded: set[str],
    path: Path,
    preferred_name: str = "",
) -> str:
    candidates = []
    excluded_norm = {_normal_header(name) for name in excluded}
    preferred_norm = _normal_header(preferred_name)
    for name in fieldnames:
        if not _clean_header(name) or _normal_header(name) in excluded_norm:
            continue
        values = [row.get(name, "") for row in rows]
        non_empty = [value for value in values if str(value).strip()]
        if not non_empty:
            continue
        numeric_count = 0
        for value in non_empty:
            try:
                float(str(value).replace(",", ""))
                numeric_count += 1
            except ValueError:
                pass
        if numeric_count:
            name_norm = _normal_header(name)
            preferred_score = 0
            if preferred_norm:
                if name_norm == preferred_norm:
                    preferred_score = 4
                elif name_norm == f"{preferred_norm}demand":
                    preferred_score = 3
                elif preferred_norm in name_norm:
                    preferred_score = 2
            candidates.append((preferred_score, numeric_count, name))
    if not candidates:
        raise ProfileGenerationError(f"{path.name} must contain one numeric value column.")
    candidates.sort(key=lambda item: (-item[0], -item[1], item[2]))
    return candidates[0][2]


def _profile_dates_with_optional_omitted_leap_day(
    start: dt.date, end: dt.date, periods_per_day: int, row_count: int
) -> tuple[list[dt.date], dt.date | None]:
    dates = date_range(start, end)
    if len(dates) not in (365, 366):
        raise ProfileGenerationError(
            "Base profile date range must cover one financial/calendar year "
            "(365 or 366 days)."
        )
    if row_count == len(dates) * periods_per_day:
        return dates, None

    leap_year = start.year if start.month <= 2 else end.year
    try:
        leap_day = dt.date(leap_year, 2, 29)
    except ValueError:
        leap_day = None
    if (
        leap_day is not None
        and leap_day in dates
        and row_count == (len(dates) - 1) * periods_per_day
    ):
        return [date_value for date_value in dates if date_value != leap_day], leap_day

    raise ProfileGenerationError(
        f"Base profile has {row_count:,} usable rows, which does not match "
        f"{len(dates)} days and {periods_per_day} periods/day."
    )


def _load_sequential_base_profile(
    path: Path,
    fieldnames: list[str],
    rows: list[dict[str, str]],
    base_start: dt.date,
    base_end: dt.date,
) -> tuple[list[BaseDay], int, str]:
    value_col = _find_value_column(fieldnames, rows, set(), path)
    values = []
    for line_number, row in enumerate(rows, start=2):
        raw_value = row.get(value_col, "")
        if not str(raw_value).strip():
            continue
        values.append(_as_float(raw_value, f"{path.name} line {line_number} {value_col}"))
    if not values:
        raise ProfileGenerationError(f"{path.name} contains no profile values.")

    dates = date_range(base_start, base_end)
    periods_per_day = None
    profile_dates = None
    for candidate_periods in (96, 24):
        try:
            candidate_dates, _ = _profile_dates_with_optional_omitted_leap_day(
                base_start, base_end, candidate_periods, len(values)
            )
        except ProfileGenerationError:
            continue
        periods_per_day = candidate_periods
        profile_dates = candidate_dates
        break
    if periods_per_day is None or profile_dates is None:
        try:
            _profile_dates_with_optional_omitted_leap_day(
                base_start, base_end, 48, len(values)
            )
        except ProfileGenerationError:
            pass
        else:
            raise ProfileGenerationError(
                "48-period demand profiles are no longer supported. Use 24 or 96 periods/day."
            )
        expected = ", ".join(
            f"{len(dates) * periods:,}" for periods in (24, 96)
        )
        raise ProfileGenerationError(
            f"{path.name} has {len(values):,} usable rows. Expected one of: {expected}."
        )

    base_days = []
    index = 0
    for profile_date in profile_dates:
        day_values = values[index : index + periods_per_day]
        if len(day_values) != periods_per_day:
            raise ProfileGenerationError(f"{path.name} ends mid-day.")
        base_days.append(BaseDay(profile_date, tuple(day_values)))
        index += periods_per_day
    return base_days, periods_per_day, value_col


def load_base_profile(
    path: Path | str, base_start: dt.date, base_end: dt.date
) -> tuple[list[BaseDay], int, str]:
    path = Path(path)
    fieldnames, rows = _read_tabular_rows(path)
    try:
        month_col = _find_required_column(fieldnames, "Month", path)
        day_col = _find_required_column(fieldnames, "Day", path)
        period_col = _find_required_column(fieldnames, "Period", path)
    except ProfileGenerationError:
        return _load_sequential_base_profile(path, fieldnames, rows, base_start, base_end)

    value_col = _find_value_column(
        fieldnames, rows, {month_col, day_col, period_col}, path
    )

    by_key: dict[tuple[int, int, int], float] = {}
    periods_seen: set[int] = set()
    usable_row_count = 0
    for line_number, row in enumerate(rows, start=2):
        raw_value = row.get(value_col, "")
        if not str(raw_value).strip():
            continue
        month = _as_int(row.get(month_col, ""), f"{path.name} line {line_number} Month")
        day = _as_int(row.get(day_col, ""), f"{path.name} line {line_number} Day")
        period = _as_int(
            row.get(period_col, ""), f"{path.name} line {line_number} Period"
        )
        value = _as_float(raw_value, f"{path.name} line {line_number} {value_col}")
        key = (month, day, period)
        if key in by_key:
            raise ProfileGenerationError(
                f"{path.name} contains duplicate Month/Day/Period row {key}."
            )
        by_key[key] = value
        periods_seen.add(period)
        usable_row_count += 1

    if not periods_seen:
        raise ProfileGenerationError(f"{path.name} contains no profile rows.")
    periods_per_day = max(periods_seen)
    if periods_per_day not in (24, 96):
        raise ProfileGenerationError(
            f"Unsupported periods per day: {periods_per_day}. "
            "Use 24 or 96 period profiles; 48-period profiles are no longer supported."
        )
    expected_periods = set(range(1, periods_per_day + 1))
    if periods_seen != expected_periods:
        raise ProfileGenerationError(
            f"{path.name} Period values must be consecutive from 1 to {periods_per_day}."
        )

    dates, omitted_leap_day = _profile_dates_with_optional_omitted_leap_day(
        base_start, base_end, periods_per_day, usable_row_count
    )

    base_days = []
    missing = []
    for profile_date in date_range(base_start, base_end):
        if profile_date == omitted_leap_day:
            continue
        values = []
        for period in range(1, periods_per_day + 1):
            key = (profile_date.month, profile_date.day, period)
            if key not in by_key:
                missing.append(key)
                continue
            values.append(by_key[key])
        if len(values) == periods_per_day:
            base_days.append(BaseDay(profile_date, tuple(values)))
    if missing:
        preview = ", ".join(str(key) for key in missing[:5])
        raise ProfileGenerationError(
            f"{path.name} is missing expected Month/Day/Period rows: {preview}."
        )
    return base_days, periods_per_day, value_col


def load_targets(path: Path | str, preferred_name: str = "") -> dict[int, float]:
    path = Path(path)
    fieldnames, rows = _read_tabular_rows(path)
    date_col = _find_required_column(fieldnames, "DateTime", path)
    value_col = _find_value_column(
        fieldnames, rows, {date_col}, path, preferred_name=preferred_name
    )
    targets = {}
    for line_number, row in enumerate(rows, start=2):
        date_value = parse_date(row.get(date_col, ""))
        value = _as_float(row.get(value_col, ""), f"{path.name} line {line_number}")
        targets[date_value.year] = value
    if not targets:
        raise ProfileGenerationError(f"{path.name} contains no target rows.")
    return targets


def _find_alias_column(
    fieldnames: Iterable[str], aliases: Iterable[str], path: Path, label: str
) -> str:
    by_normal_name = {_normal_header(name): name for name in fieldnames}
    for alias in aliases:
        match = by_normal_name.get(_normal_header(alias))
        if match:
            return match
    raise ProfileGenerationError(
        f"{path.name} must contain a '{label}' column."
    )


def parse_financial_year_start(value: str) -> int:
    text = str(value).strip()
    match = re.fullmatch(r"(\d{4})(?:\s*-\s*(\d{2}|\d{4}))?", text)
    if not match:
        raise ProfileGenerationError(
            f"Could not parse financial year '{value}'. Use YYYY-YY."
        )
    start_year = int(match.group(1))
    suffix = match.group(2)
    if suffix:
        expected_year = start_year + 1
        expected = str(expected_year) if len(suffix) == 4 else str(expected_year % 100).zfill(2)
        if suffix != expected:
            raise ProfileGenerationError(
                f"Financial year '{value}' is inconsistent; expected {fiscal_year_label(start_year)}."
            )
    return start_year


def load_rooftop_trajectory(path: Path | str) -> dict[int, float]:
    path = Path(path)
    fieldnames, rows = _read_tabular_rows(path)
    year_col = _find_alias_column(
        fieldnames,
        ("Financial Year", "FY", "Projection Year", "Year"),
        path,
        "Financial Year",
    )
    capacity_col = _find_alias_column(
        fieldnames,
        (
            "Cumulative MW",
            "Cumulative Capacity MW",
            "Rooftop Capacity MW",
            "Installed Capacity MW",
            "Capacity MW",
        ),
        path,
        "Cumulative MW",
    )
    trajectory: dict[int, float] = {}
    for line_number, row in enumerate(rows, start=2):
        year = parse_financial_year_start(row.get(year_col, ""))
        capacity = _as_float(
            row.get(capacity_col, ""), f"{path.name} line {line_number} {capacity_col}"
        )
        if not math.isfinite(capacity):
            raise ProfileGenerationError(
                f"{path.name} line {line_number} cumulative capacity must be finite."
            )
        if capacity < 0:
            raise ProfileGenerationError(
                f"{path.name} line {line_number} cumulative capacity cannot be negative."
            )
        if year in trajectory:
            raise ProfileGenerationError(
                f"{path.name} contains duplicate financial year {fiscal_year_label(year)}."
            )
        trajectory[year] = capacity

    if not trajectory:
        raise ProfileGenerationError(f"{path.name} contains no rooftop trajectory rows.")
    previous_capacity: float | None = None
    for year in sorted(trajectory):
        capacity = trajectory[year]
        if previous_capacity is not None and capacity < previous_capacity:
            raise ProfileGenerationError(
                f"{path.name} cumulative capacity must be non-decreasing; "
                f"{fiscal_year_label(year)} is lower than the preceding milestone."
            )
        previous_capacity = capacity
    return trajectory


def validate_rooftop_trajectory_coverage(
    trajectory: dict[int, float],
    base_financial_year: int,
    projection_start_year: int,
    projection_end_year: int,
) -> None:
    if projection_start_year <= base_financial_year:
        raise ProfileGenerationError(
            "Rooftop-adjusted projections must begin after the reference financial year, "
            "because the rooftop baseline is measured at that financial-year end."
        )
    if base_financial_year not in trajectory:
        raise ProfileGenerationError(
            "Rooftop trajectory must include the reference financial-year-end "
            f"baseline for {fiscal_year_label(base_financial_year)}."
        )
    first_year = min(trajectory)
    last_year = max(trajectory)
    if first_year > base_financial_year or projection_start_year < base_financial_year:
        raise ProfileGenerationError("Rooftop trajectory does not cover the projection start.")
    if last_year < projection_end_year:
        raise ProfileGenerationError(
            "Rooftop trajectory must include a milestone covering the final projection "
            f"financial year {fiscal_year_label(projection_end_year)}; extrapolation is not allowed."
        )


def _profile_period_count(periods_seen: set[int], path: Path) -> int:
    if not periods_seen:
        raise ProfileGenerationError(f"{path.name} contains no rooftop profile rows.")
    periods_per_day = max(periods_seen)
    if periods_per_day not in (24, 96):
        raise ProfileGenerationError(
            f"Unsupported rooftop periods per day: {periods_per_day}. Use 24 or 96."
        )
    if periods_seen != set(range(1, periods_per_day + 1)):
        raise ProfileGenerationError(
            f"{path.name} Period values must be consecutive from 1 to {periods_per_day}."
        )
    return periods_per_day


def load_rooftop_profile(path: Path | str, mode: str) -> RooftopProfile:
    path = Path(path)
    mode = str(mode).strip().lower()
    if mode not in {"daily", "monthly", "annual"}:
        raise ProfileGenerationError("Rooftop profile mode must be Daily, Monthly, or Annual.")

    fieldnames, rows = _read_tabular_rows(path)
    period_col = _find_alias_column(
        fieldnames, ("Period", "Time Block", "Block"), path, "Period"
    )
    cf_col = _find_alias_column(
        fieldnames,
        (
            "Rooftop CF",
            "CF",
            "Capacity Factor",
            "PU",
            "Per Unit",
            "Normalized Generation",
        ),
        path,
        "Rooftop CF",
    )
    month_col = None
    day_col = None
    if mode in {"monthly", "annual"}:
        month_col = _find_alias_column(fieldnames, ("Month",), path, "Month")
    if mode == "annual":
        day_col = _find_alias_column(fieldnames, ("Day",), path, "Day")

    values: dict[tuple[int, ...], float] = {}
    periods_seen: set[int] = set()
    for line_number, row in enumerate(rows, start=2):
        period = _as_int(
            row.get(period_col, ""), f"{path.name} line {line_number} Period"
        )
        cf = _as_float(row.get(cf_col, ""), f"{path.name} line {line_number} {cf_col}")
        if not 0 <= cf <= 1:
            raise ProfileGenerationError(
                f"{path.name} line {line_number} Rooftop CF must be between 0 and 1."
            )
        if mode == "daily":
            key = (period,)
        else:
            month = _as_int(
                row.get(month_col, ""), f"{path.name} line {line_number} Month"
            )
            if not 1 <= month <= 12:
                raise ProfileGenerationError(
                    f"{path.name} line {line_number} Month must be between 1 and 12."
                )
            if mode == "monthly":
                key = (month, period)
            else:
                day = _as_int(
                    row.get(day_col, ""), f"{path.name} line {line_number} Day"
                )
                try:
                    dt.date(2024 if month == 2 and day == 29 else 2023, month, day)
                except ValueError as exc:
                    raise ProfileGenerationError(
                        f"{path.name} line {line_number} has invalid Month/Day {month}/{day}."
                    ) from exc
                key = (month, day, period)
        if key in values:
            raise ProfileGenerationError(
                f"{path.name} contains duplicate rooftop profile row {key}."
            )
        values[key] = cf
        periods_seen.add(period)

    periods_per_day = _profile_period_count(periods_seen, path)
    if mode == "daily":
        expected = {(period,) for period in range(1, periods_per_day + 1)}
        allowed_expected = (expected,)
    elif mode == "monthly":
        expected = {
            (month, period)
            for month in range(1, 13)
            for period in range(1, periods_per_day + 1)
        }
        allowed_expected = (expected,)
    else:
        allowed_expected = tuple(
            {
                (date_value.month, date_value.day, period)
                for date_value in date_range(dt.date(year, 1, 1), dt.date(year, 12, 31))
                for period in range(1, periods_per_day + 1)
            }
            for year in (2023, 2024)
        )
    if not any(set(values) == expected_keys for expected_keys in allowed_expected):
        closest = min(allowed_expected, key=lambda keys: len(set(values) ^ keys))
        missing = sorted(closest - set(values))
        extra = sorted(set(values) - closest)
        details = []
        if missing:
            details.append(f"missing {', '.join(map(str, missing[:5]))}")
        if extra:
            details.append(f"unexpected {', '.join(map(str, extra[:5]))}")
        raise ProfileGenerationError(
            f"{path.name} does not contain a complete {mode} rooftop profile: "
            + "; ".join(details)
            + "."
        )

    return RooftopProfile(
        mode=mode,
        periods_per_day=periods_per_day,
        values=values,
        template_cuf=sum(values.values()) / len(values),
        row_count=len(values),
    )


def expand_rooftop_profile(
    profile: RooftopProfile, projection_dates: list[dt.date]
) -> list[float]:
    expanded: list[float] = []
    for date_value in projection_dates:
        for period in range(1, profile.periods_per_day + 1):
            if profile.mode == "daily":
                value = profile.values[(period,)]
            elif profile.mode == "monthly":
                value = profile.values[(date_value.month, period)]
            else:
                key = (date_value.month, date_value.day, period)
                if key in profile.values:
                    value = profile.values[key]
                elif date_value.month == 2 and date_value.day == 29:
                    value = (
                        profile.values[(2, 28, period)]
                        + profile.values[(3, 1, period)]
                    ) / 2.0
                else:
                    raise ProfileGenerationError(
                        f"Annual rooftop profile is missing {key}."
                    )
            expanded.append(value)
    return expanded


def _trajectory_milestones(
    trajectory: dict[int, float],
) -> list[tuple[dt.datetime, float]]:
    return [
        (dt.datetime(year + 1, 4, 1), trajectory[year])
        for year in sorted(trajectory)
    ]


def interpolate_rooftop_capacity(
    trajectory: dict[int, float], timestamp: dt.datetime
) -> float:
    return _interpolate_rooftop_milestones(_trajectory_milestones(trajectory), timestamp)


def _interpolate_rooftop_milestones(
    milestones: list[tuple[dt.datetime, float]], timestamp: dt.datetime
) -> float:
    if timestamp < milestones[0][0] or timestamp > milestones[-1][0]:
        raise ProfileGenerationError(
            "Rooftop trajectory does not cover all projection intervals; extrapolation is not allowed."
        )
    for index, (right_time, right_capacity) in enumerate(milestones):
        if timestamp == right_time or index == 0:
            if timestamp == right_time:
                return right_capacity
            continue
        left_time, left_capacity = milestones[index - 1]
        if left_time <= timestamp <= right_time:
            elapsed = (timestamp - left_time).total_seconds()
            span = (right_time - left_time).total_seconds()
            return left_capacity + (right_capacity - left_capacity) * elapsed / span
    raise ProfileGenerationError("Could not interpolate rooftop capacity.")


def _replace_year(month: int, day: int, year: int) -> dt.date:
    try:
        return dt.date(year, month, day)
    except ValueError as exc:
        raise ProfileGenerationError(
            f"Cannot construct date for {year:04d}-{month:02d}-{day:02d}."
        ) from exc


def projection_dates_for_year(
    base_start: dt.date, base_end: dt.date, projection_start_year: int
) -> list[dt.date]:
    start = _replace_year(base_start.month, base_start.day, projection_start_year)
    end_year = projection_start_year
    if (base_end.month, base_end.day) < (base_start.month, base_start.day):
        end_year += 1
    end = _replace_year(base_end.month, base_end.day, end_year)
    return date_range(start, end)


def _date_from_month_day(year: int, month: int, day: int) -> dt.date:
    try:
        return dt.date(year, month, day)
    except ValueError:
        return dt.date(year, month, day - 1) + dt.timedelta(days=1)


def _first_weekday_on_or_after(date_value: dt.date, weekday: int) -> dt.date:
    return date_value + dt.timedelta(days=(weekday - date_value.weekday()) % 7)


def _base_calendar_year_for_date(
    target_date: dt.date, base_start: dt.date, base_end: dt.date
) -> int:
    if (target_date.month, target_date.day) >= (base_start.month, base_start.day):
        return base_start.year
    return base_end.year


def _calendar_shift_for_date(
    target_date: dt.date,
    base_start: dt.date,
    base_end: dt.date,
    omitted_base_leap_day: bool = False,
) -> int:
    base_year = _base_calendar_year_for_date(target_date, base_start, base_end)
    base_equivalent = _date_from_month_day(
        base_year, target_date.month, target_date.day
    )
    raw_shift = (target_date.weekday() - base_equivalent.weekday()) % 7
    shift = raw_shift if raw_shift <= 3 else raw_shift - 7

    if (
        omitted_base_leap_day
        and target_date.month == 3
        and target_date.year % 4 != 0
        and raw_shift == 3
    ):
        return -4

    # Use the previous matching weekday when a three-day forward shift would
    # move the first calendar segment too far into the following week.
    fy_start_year = (
        target_date.year
        if (target_date.month, target_date.day) >= (base_start.month, base_start.day)
        else target_date.year - 1
    )
    fy_start = dt.date(fy_start_year, base_start.month, base_start.day)
    first_full_week = _first_weekday_on_or_after(fy_start, base_start.weekday())
    leading_days = (first_full_week - fy_start).days
    if (
        raw_shift == 3
        and (target_date.month, target_date.day) >= (base_start.month, base_start.day)
        and base_start.year % 4 == 0
        and fy_start.year % 4 != 0
        and leading_days >= 4
        and target_date >= first_full_week
    ):
        return -4

    # Leap day in the projected Jan-Mar segment pushes March one weekday
    # further forward when the base Jan-Mar segment is non-leap.
    if (
        raw_shift == 4
        and target_date.month == 3
        and target_date.year % 4 == 0
        and not (base_end.year % 4 == 0)
    ):
        return 4

    return shift


def _is_in_calendar_segment(
    date_value: dt.date, base_start: dt.date, base_end: dt.date, month: int
) -> bool:
    if month >= base_start.month:
        return date_value.year == base_start.year and date_value.month >= base_start.month
    return date_value.year == base_end.year and date_value.month <= base_end.month


def _projection_fy_start(target_date: dt.date, base_start: dt.date) -> dt.date:
    year = target_date.year
    if (target_date.month, target_date.day) < (base_start.month, base_start.day):
        year -= 1
    return dt.date(year, base_start.month, base_start.day)


def _projection_fy_end(fy_start: dt.date, base_start: dt.date, base_end: dt.date) -> dt.date:
    end_year = fy_start.year
    if (base_end.month, base_end.day) < (base_start.month, base_start.day):
        end_year += 1
    return dt.date(end_year, base_end.month, base_end.day)


def _map_source_date(
    target_date: dt.date,
    base_start: dt.date,
    base_end: dt.date,
    omitted_base_leap_day: bool = False,
) -> dt.date:
    base_year = _base_calendar_year_for_date(target_date, base_start, base_end)
    base_equivalent = _date_from_month_day(
        base_year, target_date.month, target_date.day
    )
    shift = _calendar_shift_for_date(
        target_date, base_start, base_end, omitted_base_leap_day
    )
    source_date = base_equivalent + dt.timedelta(days=shift)

    # Keep Dec/Jan boundary adjustments in their own calendar segment before
    # applying the fiscal-year edge behavior.
    if (
        target_date.month >= base_start.month
        and source_date.year == base_end.year
        and source_date.month == 1
    ):
        while source_date.year == base_end.year and source_date.month == 1:
            source_date -= dt.timedelta(days=7)
    if target_date.month < base_start.month and source_date < dt.date(base_end.year, 1, 1):
        while source_date < dt.date(base_end.year, 1, 1):
            source_date += dt.timedelta(days=7)

    fy_start = _projection_fy_start(target_date, base_start)
    first_full_week = _first_weekday_on_or_after(fy_start, base_start.weekday())
    leading_days = (first_full_week - fy_start).days

    if target_date.month >= base_start.month and source_date < base_start:
        if target_date < first_full_week and leading_days < 4:
            fy_end = _projection_fy_end(fy_start, base_start, base_end)
            fy_day_count = (fy_end - fy_start).days + 1
            use_segment_start = False
            if omitted_base_leap_day and fy_day_count == 365 and leading_days == 2:
                use_segment_start = True
            if base_start.weekday() in (4, 5) and target_date != fy_start and leading_days == 3:
                use_segment_start = True
            if use_segment_start:
                while source_date < base_start:
                    source_date += dt.timedelta(days=7)
            else:
                while source_date < base_start:
                    source_date += dt.timedelta(days=364)
        else:
            while source_date < base_start:
                source_date += dt.timedelta(days=7)

    fy_end = _projection_fy_end(fy_start, base_start, base_end)
    fy_day_count = (fy_end - fy_start).days + 1
    if (
        target_date == fy_start
        and shift == 0
        and target_date.weekday() == base_start.weekday()
        and fy_day_count == 365
        and not omitted_base_leap_day
    ):
        source_date = base_end

    if (
        omitted_base_leap_day
        and target_date == fy_start
        and shift == 1
        and source_date + dt.timedelta(days=364) <= base_end
    ):
        source_date += dt.timedelta(days=364)

    if target_date.month < base_start.month and source_date > base_end:
        if omitted_base_leap_day and shift in (1, 2) and target_date != fy_end:
            while source_date > base_end:
                source_date -= dt.timedelta(days=7)
        elif shift in (1, 2):
            while source_date > base_end:
                source_date -= dt.timedelta(days=364)
        else:
            while source_date > base_end:
                source_date -= dt.timedelta(days=7)

    if (
        target_date == fy_end
        and fy_day_count == 366
        and base_start.weekday() == 4
        and source_date == base_end
        and target_date.weekday() == base_end.weekday()
    ):
        source_date = base_start

    if source_date < base_start or source_date > base_end:
        raise ProfileGenerationError(
            f"Could not map {target_date.isoformat()} into the base profile."
        )
    return source_date


def _synthetic_missing_base_day(
    source_date: dt.date, by_date: dict[dt.date, BaseDay], periods_per_day: int
) -> BaseDay:
    earlier = [date_value for date_value in by_date if date_value < source_date]
    later = [date_value for date_value in by_date if date_value > source_date]
    if earlier and later:
        left = by_date[max(earlier)]
        right = by_date[min(later)]
        level = (sum(left.values) / periods_per_day + sum(right.values) / periods_per_day) / 2.0
    elif earlier:
        left = by_date[max(earlier)]
        level = sum(left.values) / periods_per_day
    elif later:
        right = by_date[min(later)]
        level = sum(right.values) / periods_per_day
    else:
        raise ProfileGenerationError("Cannot synthesize a missing base day.")
    return BaseDay(source_date, tuple(level for _ in range(periods_per_day)))


def map_projection_days_by_calendar(
    base_days: list[BaseDay],
    projection_dates: list[dt.date],
    base_start: dt.date,
    base_end: dt.date,
) -> list[BaseDay]:
    by_date = {day.date: day for day in base_days}
    periods_per_day = len(base_days[0].values) if base_days else 0
    omitted_leap_day = None
    try:
        candidate = dt.date(base_end.year, 2, 29)
    except ValueError:
        candidate = None
    if candidate is not None and base_start <= candidate <= base_end:
        if candidate not in by_date:
            omitted_leap_day = candidate
    mapped = []
    for target_date in projection_dates:
        source_date = _map_source_date(
            target_date,
            base_start,
            base_end,
            omitted_base_leap_day=omitted_leap_day is not None,
        )
        if source_date in by_date:
            mapped.append(by_date[source_date])
            continue
        if source_date.month == 2 and source_date.day == 29 and periods_per_day:
            mapped.append(
                _synthetic_missing_base_day(source_date, by_date, periods_per_day)
            )
            continue
        raise ProfileGenerationError(
            f"Base profile does not contain mapped source date {source_date}."
        )
    return mapped


def map_projection_days(
    base_days: list[BaseDay],
    projection_dates: list[dt.date],
    base_start: dt.date,
    base_end: dt.date,
) -> list[BaseDay]:
    return map_projection_days_by_calendar(
        base_days, projection_dates, base_start, base_end
    )


def _rank_quantiles(values: list[float]) -> list[float]:
    if len(values) == 1:
        return [1.0]
    order = sorted(range(len(values)), key=lambda index: (values[index], index))
    denominator = len(values) - 1
    quantiles = [0.0] * len(values)
    for rank, index in enumerate(order):
        quantiles[index] = rank / denominator
    return quantiles


def scale_year(
    mapped_values: list[float],
    target_peak_mw: float,
    target_energy_gwh: float,
    periods_per_day: int,
) -> tuple[list[int], float, float]:
    if target_peak_mw <= 0:
        raise ProfileGenerationError("Target peak must be positive.")
    if target_energy_gwh <= 0:
        raise ProfileGenerationError("Target energy must be positive.")
    mapped_peak = max(mapped_values)
    if mapped_peak <= 0:
        raise ProfileGenerationError("Mapped base peak must be positive.")

    interval_hours = 24.0 / periods_per_day
    target_mw_sum = target_energy_gwh * 1000.0 / interval_hours
    q_values = _rank_quantiles(mapped_values)
    peak_growth = target_peak_mw / mapped_peak
    denominator = sum(value * (1.0 - q) for value, q in zip(mapped_values, q_values))
    if denominator <= 0:
        raise ProfileGenerationError("Cannot solve low-rank growth rate.")
    numerator = target_mw_sum - sum(
        value * peak_growth * q for value, q in zip(mapped_values, q_values)
    )
    low_rank_growth = numerator / denominator

    projected = []
    for value, q in zip(mapped_values, q_values):
        growth = low_rank_growth * (1.0 - q) + peak_growth * q
        projected_value = int(round(value * growth))
        if projected_value < 0:
            raise ProfileGenerationError(
                "Projection would create negative demand. Check peak and energy targets."
            )
        projected.append(projected_value)
    return projected, peak_growth, low_rank_growth


def sanitize_name(value: str, default: str = "projected_profile") -> str:
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", value.strip())
    value = value.strip("._-")
    return value or default


def unique_output_path(output_dir: Path | str, filename: str) -> Path:
    output_dir = Path(output_dir)
    candidate = output_dir / filename
    if not candidate.exists():
        return candidate
    stem = candidate.stem
    suffix = candidate.suffix
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    return output_dir / f"{stem}_{stamp}{suffix}"


def _normalize_for_graph(values: list[float] | list[int], label: str) -> list[float]:
    peak = max(values) if values else 0
    if peak <= 0:
        raise ProfileGenerationError(f"{label} peak must be positive for charting.")
    return [round(float(value) / float(peak), 6) for value in values]


def _profile_looks_normalized(values: list[float]) -> bool:
    if not values:
        return False
    return min(values) >= 0 and max(values) <= 1.05


def generate_profiles(
    base_profile_path: Path | str,
    peak_projection_path: Path | str,
    energy_projection_path: Path | str,
    base_start_date: dt.date,
    base_end_date: dt.date,
    projection_start_year: int,
    projection_end_year: int,
    output_dir: Path | str,
    profile_name: str = "",
    rooftop_profile_path: Path | str | None = None,
    rooftop_trajectory_path: Path | str | None = None,
    rooftop_profile_mode: str | None = None,
) -> GenerationResult:
    if projection_end_year < projection_start_year:
        raise ProfileGenerationError(
            "Projection end year must be greater than or equal to start year."
        )
    validate_base_financial_year_period(base_start_date, base_end_date)
    output_dir = Path(output_dir)
    if not output_dir.exists() or not output_dir.is_dir():
        raise ProfileGenerationError("Output folder does not exist.")

    base_days, periods_per_day, _ = load_base_profile(
        base_profile_path, base_start_date, base_end_date
    )
    interval_hours = 24.0 / periods_per_day
    base_values = [value for base_day in base_days for value in base_day.values]
    base_profile_peak = float(max(base_values))
    base_profile_energy = sum(base_values) * interval_hours / 1000.0
    base_profile_is_normalized = _profile_looks_normalized(base_values)
    profile_slug = sanitize_name(profile_name)
    peak_targets = load_targets(peak_projection_path, preferred_name=profile_name)
    energy_targets = load_targets(energy_projection_path, preferred_name=profile_name)
    rooftop_args = (
        rooftop_profile_path,
        rooftop_trajectory_path,
        rooftop_profile_mode,
    )
    rooftop_enabled = any(value is not None and value != "" for value in rooftop_args)
    if rooftop_enabled and not all(value is not None and value != "" for value in rooftop_args):
        raise ProfileGenerationError(
            "Rooftop mode requires a profile, trajectory, and profile mode."
        )

    rooftop_profile = None
    rooftop_trajectory = None
    rooftop_milestones = None
    baseline_rooftop_capacity = None
    if rooftop_enabled:
        rooftop_profile = load_rooftop_profile(
            rooftop_profile_path, str(rooftop_profile_mode)
        )
        if rooftop_profile.periods_per_day != periods_per_day:
            raise ProfileGenerationError(
                "Rooftop profile resolution must exactly match the demand profile "
                f"({rooftop_profile.periods_per_day} versus {periods_per_day} periods/day)."
            )
        rooftop_trajectory = load_rooftop_trajectory(rooftop_trajectory_path)
        validate_rooftop_trajectory_coverage(
            rooftop_trajectory,
            base_start_date.year,
            projection_start_year,
            projection_end_year,
        )
        baseline_rooftop_capacity = rooftop_trajectory[base_start_date.year]
        rooftop_milestones = _trajectory_milestones(rooftop_trajectory)

    for year in range(projection_start_year, projection_end_year + 1):
        if year not in peak_targets:
            raise ProfileGenerationError(f"Peak projection is missing target year {year}.")
        if year not in energy_targets:
            raise ProfileGenerationError(f"Energy projection is missing target year {year}.")

    if rooftop_enabled:
        filename = f"{profile_slug}_rooftop_adjusted.csv"
        output_columns = ROOFTOP_OUTPUT_COLUMNS
    else:
        filename = (
            f"{profile_slug}_projected_demand_"
            f"{projection_start_year}_{projection_end_year}.csv"
        )
        output_columns = OUTPUT_COLUMNS
    output_path = unique_output_path(output_dir, filename)

    summaries = []
    graph_series = []
    rows_written = 0
    with output_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=output_columns)
        writer.writeheader()

        for year in range(projection_start_year, projection_end_year + 1):
            projection_dates = projection_dates_for_year(
                base_start_date, base_end_date, year
            )
            mapped_days = map_projection_days(
                base_days,
                projection_dates,
                base_start_date,
                base_end_date,
            )
            mapped_values = [
                value for mapped_day in mapped_days for value in mapped_day.values
            ]
            projected, peak_growth, low_rank_growth = scale_year(
                mapped_values, peak_targets[year], energy_targets[year], periods_per_day
            )

            adjusted_values: list[float] | None = None
            rooftop_generation: list[float] | None = None
            incremental_capacities: list[float] | None = None
            effective_cuf: float | None = None
            if rooftop_enabled:
                assert rooftop_profile is not None
                assert rooftop_trajectory is not None
                assert rooftop_milestones is not None
                assert baseline_rooftop_capacity is not None
                rooftop_cf_values = expand_rooftop_profile(
                    rooftop_profile, projection_dates
                )
                incremental_capacities = []
                rooftop_generation = []
                for index, projection_date in enumerate(projection_dates):
                    for period_index in range(periods_per_day):
                        timestamp = dt.datetime.combine(
                            projection_date, dt.time()
                        ) + dt.timedelta(hours=period_index * interval_hours)
                        total_capacity = _interpolate_rooftop_milestones(
                            rooftop_milestones, timestamp
                        )
                        incremental_capacity = total_capacity - baseline_rooftop_capacity
                        if incremental_capacity < -1e-9:
                            raise ProfileGenerationError(
                                "Interpolated rooftop capacity falls below the reference baseline."
                            )
                        incremental_capacity = max(0.0, incremental_capacity)
                        value_index = index * periods_per_day + period_index
                        incremental_capacities.append(incremental_capacity)
                        rooftop_generation.append(
                            incremental_capacity * rooftop_cf_values[value_index]
                        )
                adjusted_values = [
                    float(before) - generation
                    for before, generation in zip(projected, rooftop_generation)
                ]
                available_capacity_hours = (
                    sum(incremental_capacities) * interval_hours
                )
                if available_capacity_hours > 0:
                    effective_cuf = (
                        sum(rooftop_generation) * interval_hours
                        / available_capacity_hours
                    )

            index = 0
            for projection_date, mapped_day in zip(projection_dates, mapped_days):
                if len(mapped_day.values) != periods_per_day:
                    raise ProfileGenerationError("Internal period-count mismatch.")
                for period in range(1, periods_per_day + 1):
                    if rooftop_enabled:
                        assert adjusted_values is not None
                        assert rooftop_generation is not None
                        assert incremental_capacities is not None
                        writer.writerow(
                            {
                                "Financial Year": fiscal_year_label(year),
                                "year": projection_date.year,
                                "month": projection_date.month,
                                "day": projection_date.day,
                                "period": period,
                                "demand before rooftop": f"{projected[index]:.3f}",
                                "incremental rooftop capacity": f"{incremental_capacities[index]:.3f}",
                                "incremental rooftop generation": f"{rooftop_generation[index]:.3f}",
                                "projected demand": f"{adjusted_values[index]:.3f}",
                            }
                        )
                    else:
                        writer.writerow(
                            {
                                "Financial Year": fiscal_year_label(year),
                                "year": projection_date.year,
                                "month": projection_date.month,
                                "day": projection_date.day,
                                "period": period,
                                "projected demand": projected[index],
                            }
                        )
                    index += 1

            final_values = adjusted_values if adjusted_values is not None else projected
            before_peak = float(max(projected))
            final_peak = float(max(final_values))
            final_energy = sum(final_values) * interval_hours / 1000.0
            rooftop_generation_gwh = (
                None
                if rooftop_generation is None
                else sum(rooftop_generation) * interval_hours / 1000.0
            )
            peak_index = max(range(len(final_values)), key=final_values.__getitem__)
            peak_date = projection_dates[peak_index // periods_per_day]
            peak_period = peak_index % periods_per_day + 1
            year_end_capacity = (
                None
                if rooftop_trajectory is None
                else _interpolate_rooftop_milestones(
                    rooftop_milestones, dt.datetime(year + 1, 4, 1)
                )
            )

            summaries.append(
                YearSummary(
                    projection_year=year,
                    row_count=len(projected),
                    day_count=len(projection_dates),
                    periods_per_day=periods_per_day,
                    base_profile_peak_mw=base_profile_peak,
                    base_profile_energy_gwh=base_profile_energy,
                    base_profile_is_normalized=base_profile_is_normalized,
                    target_peak_mw=peak_targets[year],
                    achieved_peak_mw=final_peak,
                    target_energy_gwh=energy_targets[year],
                    achieved_energy_gwh=final_energy,
                    mapped_base_peak_mw=float(max(mapped_values)),
                    mapped_base_energy_gwh=sum(mapped_values) * interval_hours / 1000.0,
                    peak_growth_rate=peak_growth,
                    low_rank_growth_rate=low_rank_growth,
                    rooftop_enabled=rooftop_enabled,
                    before_rooftop_peak_mw=before_peak if rooftop_enabled else None,
                    before_rooftop_energy_gwh=(
                        sum(projected) * interval_hours / 1000.0
                        if rooftop_enabled
                        else None
                    ),
                    peak_reduction_mw=(
                        before_peak - final_peak if rooftop_enabled else None
                    ),
                    peak_reduction_percent=(
                        (before_peak - final_peak) / before_peak * 100.0
                        if rooftop_enabled and before_peak
                        else None
                    ),
                    rooftop_generation_gwh=rooftop_generation_gwh,
                    rooftop_template_cuf_percent=(
                        rooftop_profile.template_cuf * 100.0
                        if rooftop_profile is not None
                        else None
                    ),
                    rooftop_effective_cuf_percent=(
                        effective_cuf * 100.0 if effective_cuf is not None else None
                    ),
                    rooftop_year_end_capacity_mw=year_end_capacity,
                    adjusted_minimum_mw=(
                        float(min(final_values)) if rooftop_enabled else None
                    ),
                    adjusted_peak_date=(peak_date.isoformat() if rooftop_enabled else None),
                    adjusted_peak_period=(peak_period if rooftop_enabled else None),
                )
            )
            graph_year = {
                "id": f"year-{year}",
                "label": fiscal_year_label(year),
                "year": year,
                "start_date": projection_dates[0].isoformat(),
                "end_date": projection_dates[-1].isoformat(),
                "day_count": len(projection_dates),
                "periods_per_day": periods_per_day,
            }
            if rooftop_enabled:
                assert adjusted_values is not None
                assert rooftop_generation is not None
                graph_year.update(
                    {
                        "before_rooftop": {
                            "label": "Before rooftop",
                            "points": projected,
                        },
                        "adjusted": {
                            "label": "Adjusted grid demand",
                            "points": adjusted_values,
                        },
                        "rooftop_generation": {
                            "label": "Rooftop generation",
                            "points": rooftop_generation,
                        },
                    }
                )
            else:
                graph_year.update(
                    {
                        "mapped_base": {
                            "label": "Mapped base",
                            "points": _normalize_for_graph(
                                mapped_values, f"Mapped base {year}"
                            ),
                        },
                        "projected": {
                            "label": "Projected",
                            "points": _normalize_for_graph(
                                projected, f"Projection {year}"
                            ),
                        },
                    }
                )
            graph_series.append(graph_year)
            rows_written += len(projected)

    graph = {
        "normalization": "absolute_mw" if rooftop_enabled else "series_peak",
        "rooftop_enabled": rooftop_enabled,
        "periods_per_day": periods_per_day,
        "years": graph_series,
    }

    return GenerationResult(output_path, rows_written, tuple(summaries), graph)


def _growth_percent(new_value: float, old_value: float) -> float | None:
    if old_value <= 0:
        return None
    return ((new_value / old_value) - 1.0) * 100.0


def summaries_as_dicts(
    summaries: Iterable[YearSummary],
) -> list[dict[str, float | int | bool | None]]:
    items = list(summaries)
    rows = []
    for index, item in enumerate(items):
        if index == 0:
            previous_peak = (
                None if item.base_profile_is_normalized else item.base_profile_peak_mw
            )
            previous_energy = (
                None if item.base_profile_is_normalized else item.base_profile_energy_gwh
            )
        else:
            previous = items[index - 1]
            previous_peak = previous.target_peak_mw
            previous_energy = previous.target_energy_gwh

        rows.append(
            {
                "projection_year": item.projection_year,
                "financial_year": fiscal_year_label(item.projection_year),
                "row_count": item.row_count,
                "day_count": item.day_count,
                "periods_per_day": item.periods_per_day,
                "base_profile_peak_mw": item.base_profile_peak_mw,
                "base_profile_energy_gwh": item.base_profile_energy_gwh,
                "base_profile_is_normalized": item.base_profile_is_normalized,
                "target_peak_mw": item.target_peak_mw,
                "achieved_peak_mw": item.achieved_peak_mw,
                "target_energy_gwh": item.target_energy_gwh,
                "achieved_energy_gwh": item.achieved_energy_gwh,
                "mapped_base_peak_mw": item.mapped_base_peak_mw,
                "mapped_base_energy_gwh": item.mapped_base_energy_gwh,
                "peak_growth_rate": item.peak_growth_rate,
                "peak_growth_percent": (
                    None
                    if previous_peak is None
                    else _growth_percent(item.target_peak_mw, previous_peak)
                ),
                "energy_growth_percent": (
                    None
                    if previous_energy is None
                    else _growth_percent(item.target_energy_gwh, previous_energy)
                ),
                "low_rank_growth_rate": item.low_rank_growth_rate,
                "rooftop_enabled": item.rooftop_enabled,
                "before_rooftop_peak_mw": item.before_rooftop_peak_mw,
                "before_rooftop_energy_gwh": item.before_rooftop_energy_gwh,
                "peak_reduction_mw": item.peak_reduction_mw,
                "peak_reduction_percent": item.peak_reduction_percent,
                "rooftop_generation_gwh": item.rooftop_generation_gwh,
                "rooftop_template_cuf_percent": item.rooftop_template_cuf_percent,
                "rooftop_effective_cuf_percent": item.rooftop_effective_cuf_percent,
                "rooftop_year_end_capacity_mw": item.rooftop_year_end_capacity_mw,
                "adjusted_minimum_mw": item.adjusted_minimum_mw,
                "adjusted_peak_date": item.adjusted_peak_date,
                "adjusted_peak_period": item.adjusted_peak_period,
            }
        )
    return rows
